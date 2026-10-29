"""Read-only business projections; explicitly bound, manager-isolated reporting."""
from dataclasses import dataclass
from datetime import datetime,timedelta
from zoneinfo import ZoneInfo
from contextlib import contextmanager
import collections,hashlib,io,json,re,sqlite3,time,uuid
from pathlib import Path
TZ=ZoneInfo('Asia/Dubai')
@dataclass(frozen=True)
class Scope:
 company_id:str
 manager_id:str
 account_id:int
 organization:str='Организация не подтверждена'
 company_name:str=''
 allowed_products:tuple=()
 def authorize(self,actor):
  if type(actor) is not int or actor!=self.account_id:raise PermissionError('Нет доступа к чужому отчёту.')
  if not self.company_id or not self.manager_id or self.account_id<=0:raise PermissionError('Не подтверждена принадлежность отчёта.')

@contextmanager
def writer(path):
 db=sqlite3.connect(path,timeout=30)
 try:
  with db:yield db
 finally:db.close()

def initialize(path):
 with writer(path) as db:
  db.executescript('''CREATE TABLE IF NOT EXISTS report_reads (
   company_id TEXT NOT NULL,manager_id TEXT NOT NULL,account_id INTEGER NOT NULL,peer_id INTEGER NOT NULL,
   max_id INTEGER NOT NULL,observed_at INTEGER NOT NULL,PRIMARY KEY(company_id,account_id,peer_id));
   CREATE TABLE IF NOT EXISTS report_health (
   company_id TEXT NOT NULL,account_id INTEGER NOT NULL,updated_at INTEGER NOT NULL,connected INTEGER NOT NULL,
   receipts_at INTEGER,error TEXT,PRIMARY KEY(company_id,account_id));
   CREATE TABLE IF NOT EXISTS report_usage (
   company_id TEXT NOT NULL,manager_id TEXT NOT NULL,account_id INTEGER NOT NULL,id TEXT NOT NULL,
   phase TEXT NOT NULL,model TEXT NOT NULL,status TEXT NOT NULL,input_tokens INTEGER,output_tokens INTEGER,
   cached_tokens INTEGER,response_id TEXT,created_at INTEGER NOT NULL,usage_json TEXT NOT NULL,
   PRIMARY KEY(company_id,account_id,id));''')
  db.execute('''CREATE TABLE IF NOT EXISTS report_incoming (
   company_id TEXT NOT NULL,manager_id TEXT NOT NULL,account_id INTEGER NOT NULL,
   peer_id INTEGER NOT NULL,message_id INTEGER NOT NULL,sent_at INTEGER NOT NULL,
   text TEXT NOT NULL,observed_at INTEGER NOT NULL,
   PRIMARY KEY(company_id,account_id,peer_id,message_id))''')

def observe_read(path,scope,peer,max_id,at=None):
 if type(peer) is not int or peer<=0 or type(max_id) is not int or max_id<0:raise ValueError('Not a private read pointer')
 with writer(path) as db:
  db.execute('''INSERT INTO report_reads VALUES(?,?,?,?,?,?) ON CONFLICT(company_id,account_id,peer_id)
   DO UPDATE SET max_id=MAX(report_reads.max_id,excluded.max_id),observed_at=MAX(report_reads.observed_at,excluded.observed_at)
   WHERE manager_id=excluded.manager_id''',(scope.company_id,scope.manager_id,scope.account_id,peer,max_id,int(at or time.time())))

def period(value,now=None):
 now=datetime.fromtimestamp(time.time() if now is None else now,TZ);today=now.replace(hour=0,minute=0,second=0,microsecond=0)
 if value=='today':a,b=today,today+timedelta(days=1)
 elif value=='yesterday':a,b=today-timedelta(days=1),today
 elif value in {'7','30'}:a,b=today-timedelta(days=int(value)-1),today+timedelta(days=1)
 elif value=='all':a,b=datetime(2000,1,1,tzinfo=TZ),today+timedelta(days=1)
 else:
  match=re.fullmatch(r'(\d{2}\.\d{2}\.\d{4})\s*[-–]\s*(\d{2}\.\d{2}\.\d{4})',value)
  if not match:raise ValueError('Введите период: 01.10.2026–31.10.2026')
  a,b=[datetime.strptime(x,'%d.%m.%Y').replace(tzinfo=TZ) for x in match.groups()];b+=timedelta(days=1)
 if a>=b or a>now or b>today+timedelta(days=1):raise ValueError('Проверьте даты: начало не позже конца, без будущих дней.')
 return int(a.timestamp()),int(b.timestamp())

def epoch(value):
 if value is None:return None
 if isinstance(value,(int,float)):return int(value)
 try:
  d=datetime.fromisoformat(value.replace('Z','+00:00'))
  # Legacy SQLite datetime('now') timestamps are UTC, never machine-local.
  return int(d.replace(tzinfo=ZoneInfo('UTC')).timestamp()) if d.tzinfo is None else int(d.timestamp())
 except (ValueError,TypeError):return None

@contextmanager
def reader(path):
 db=sqlite3.connect(Path(path).resolve().as_uri()+'?mode=ro',uri=True,timeout=30);db.row_factory=sqlite3.Row
 try:db.execute('BEGIN');yield db
 finally:db.rollback();db.close()

def table(db,name):return bool(db.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name=?",(name,)).fetchone())
def rows(db,name):return [dict(x) for x in db.execute('SELECT * FROM '+name)] if table(db,name) else []
def link(peer,mid,username=''):
 return 'https://t.me/'+username.lstrip('@')+'/'+str(mid) if username else 'tg://openmessage?user_id='+str(peer)+'&message_id='+str(mid)

# Official standard GPT-4.1 rates verified 2026-10-29; USD per million tokens.
# Source: https://developers.openai.com/api/docs/models/gpt-4.1
PRICE_SOURCE='https://developers.openai.com/api/docs/models/gpt-4.1'
PRICE_CHECKED=1793232000

def cost(row):
 from ai_costs import estimate
 return estimate(row)

def safety_state(path):
 p=Path(path).parent/'telegram-safety.sqlite'
 if not p.exists():return dict(known=False,blocked=None,operation=None,created_at=None)
 try:
  with reader(p) as db:
   row=db.execute('SELECT operation,created_at FROM safety_incident WHERE slot=1').fetchone()
   return dict(known=True,blocked=bool(row),operation=row['operation'] if row else None,created_at=row['created_at'] if row else None)
 except (OSError,sqlite3.Error):return dict(known=False,blocked=None,operation=None,created_at=None)

def build(path,scope,*,actor,start,end,now=None,product=None):
 scope.authorize(actor);now=int(time.time() if now is None else now);safety=safety_state(path)
 if type(start) is not int or type(end) is not int or start>=end:raise ValueError('Некорректный период')
 inside=lambda value: (epoch(value) is not None and start<=epoch(value)<end)
 with reader(path) as db:
  clients={r['id']:r for r in rows(db,'clients') if r.get('manager_id')==scope.manager_id and (not scope.allowed_products or r.get('product') in scope.allowed_products) and (not product or r.get('product')==product)}
  peers={r['telegram_id']:r for r in clients.values() if type(r.get('telegram_id')) is int and r['telegram_id']>0}
  threads={r['id']:r for r in rows(db,'dialogue_threads') if r['company_id']==scope.company_id and r['account_id']==scope.account_id and r['manager_id']==scope.manager_id and (not scope.allowed_products or r['product'] in scope.allowed_products) and (not product or r['product']==product)}
  # Finance test chats with no customer row do not contribute customer KPIs.
  threads={k:v for k,v in threads.items() if v['peer_id'] in peers}
  outgoing={};incoming={};evidence={};details=[]
  for m in rows(db,'messages'):
   c=clients.get(m['client_id'])
   if c and type(c.get('telegram_id')) is int and c['telegram_id']>0 and m['status']=='sent' and m.get('telegram_message_id') and m.get('sent_at'):
    outgoing[(c['telegram_id'],m['telegram_message_id'])]=dict(peer=c['telegram_id'],message_id=m['telegram_message_id'],at=epoch(m['sent_at']),product=c['product'],text=m['text'],name=c.get('full_name',''),username=c.get('username',''))
  for m in rows(db,'dialogue_messages'):
   t=threads.get(m['thread_id'])
   if not t:continue
   c=peers[t['peer_id']];item=dict(peer=t['peer_id'],message_id=m['telegram_id'],at=epoch(m['sent_at']),product=t['product'],text=m['text'],name=c.get('full_name',''),username=c.get('username',''))
   evidence[(t['id'],m['telegram_id'])]=m
   (outgoing if m['direction']=='outgoing' else incoming)[(t['peer_id'],m['telegram_id'])]=item
  for m in rows(db,'report_incoming'):
   c=peers.get(m['peer_id'])
   if c and m['company_id']==scope.company_id and m['account_id']==scope.account_id and m['manager_id']==scope.manager_id:
    incoming[(m['peer_id'],m['message_id'])]=dict(peer=m['peer_id'],message_id=m['message_id'],at=epoch(m['sent_at']),product=c['product'],text=m['text'],name=c.get('full_name',''),username=c.get('username',''))
  reads={r['peer_id']:r['max_id'] for r in rows(db,'report_reads') if r['company_id']==scope.company_id and r['account_id']==scope.account_id and r['manager_id']==scope.manager_id}
  sent=[dict(r,read=(reads.get(r['peer'],-1)>=r['message_id'])) for r in outgoing.values() if inside(r['at'])]
  replies=[r for r in incoming.values() if inside(r['at'])]
  for direction,items in [('Исходящее',sent),('Ответ клиента',replies)]:
   for r in items:details.append(dict(r,direction=direction,url=link(r['peer'],r['message_id'],r['username'])))
  app_events=rows(db,'seller_application_events');apps=[]
  for a in rows(db,'seller_applications'):
   t=threads.get(a['thread_id'])
   if not t or not inside(a['created_at']):continue
   c=peers[t['peer_id']];source=evidence.get((t['id'],a['need_source_id']))
   ev=[r for r in app_events if r['application_id']==a['id']]
   apps.append(dict(a,peer=t['peer_id'],manager=scope.manager_id,name=c.get('full_name',''),username=c.get('username',''),need_text=source['text'] if source else 'Источник отсутствует в локальных данных',source_available=source is not None,url=link(t['peer_id'],a['need_source_id'],c.get('username','')),events=ev,transfers=[dict(n) for n in rows(db,'dialogue_notices') if n['thread_id']==t['id'] and n['event_type']=='ready']))
  notices=[r for r in rows(db,'dialogue_notices') if r['thread_id'] in threads and inside(r['created_at']) and r['event_type']=='ready']
  refused=[];pattern=re.compile(r'(?i)(не\s+(?:пишите|писать|беспокойте|беспокоить|связывайтесь)|удалите\s+мой\s+(?:номер|контакт)|не\s+интересует|не\s+нуж(?:но|ны|ен|на)|справимся\s+сами|вс[её]\s+(?:уже\s+)?проработано)')
  for r in replies:
   if pattern.search(r['text']):refused.append(dict(r,url=link(r['peer'],r['message_id'],r['username'])))
  cohorts=[r for r in rows(db,'seller_warm_cohorts') if r['thread_id'] in threads and r['status']=='active' and threads[r['thread_id']]['state']=='engaged' and threads[r['thread_id']]['mode'] not in {'copilot','off','paused_until'}]
  usage=[];seen=set()
  for u in rows(db,'report_usage'):
   if u['company_id']==scope.company_id and u['account_id']==scope.account_id and u['manager_id']==scope.manager_id and inside(u['created_at']) and (not scope.allowed_products or json.loads(u['usage_json']).get('product') in scope.allowed_products) and (not product or json.loads(u['usage_json']).get('product')==product):
    usage.append(u);seen.add(u.get('response_id') or u['id'])
  for u in rows(db,'seller_ai_usage'):
   if u['thread_id'] in threads and inside(u['created_at']) and (u.get('response_id') or u['id']) not in seen:
    usage.append(u);seen.add(u.get('response_id') or u['id'])
  usage=[dict(u,**cost(u)) for u in usage]
  known=[u for u in usage if u['usd_min'] is not None]
  costs=dict(usd_min=sum(u['usd_min'] for u in known),usd_max=sum(u['usd_max'] for u in known),priced_requests=len(known),unknown_requests=len(usage)-len(known),incomplete_tools=sum(not u['tools_complete'] for u in usage),tool_usd_known=sum(u.get('tool_usd_known',0) for u in usage),billing_complete=False,source=PRICE_SOURCE,checked_at=PRICE_CHECKED)
  sync_pending=db.execute('SELECT count(*) FROM sync_outbox').fetchone()[0] if table(db,'sync_outbox') else None
  health=next((r for r in rows(db,'report_health') if r['company_id']==scope.company_id and r['account_id']==scope.account_id),None)
  research=[r for r in rows(db,'research_runs') if r['client_id'] in clients and inside(r['created_at'])]
  warnings=['Прочтения: текущее подтверждённое состояние сообщений, отправленных за период; точная дата прочтения не восстанавливается.','Прогрев: текущее состояние, а не исторический снимок на конец выбранного периода.','Отказы: сообщения с явными фразами отказа; возможна ручная проверка. Другие формулировки автоматически не включаются.','Расходы — оценка стоимости токенов GPT-4.1 по стандартному тарифу, проверенному 29.10.2026: вход $2, кэш $0.50, выход $8 за миллион. Не счёт провайдера. Инструменты, хранилище и другие модели могут не входить. Старый кэш без детализации даёт диапазон. Источник: '+PRICE_SOURCE,'Старая история, анализ и запросы базы знаний могли выполняться без учёта usage; отсутствие записи не означает отсутствие расхода.']
  if safety['blocked']:warnings.append('Клиентские отправки остановлены защитой: '+str(safety['operation'])+' с '+stamp(safety['created_at'])+'. Отчёты продолжают работать; снятие защиты не выполняется отчётом.')
  if not health or now-health['updated_at']>120:warnings.append('Нет свежего подтверждения работы сборщика отчётов.')
  if not health or not health.get('receipts_at'):warnings.append('Начальная сверка прочтений ещё не завершена; показатель неполный.')
  return dict(scope=scope.__dict__,start=start,end=end,generated_at=now,product=product,metrics=dict(sent_messages=len(sent),contacted_clients=len({r['peer'] for r in sent}),read_messages=sum(r['read'] for r in sent),read_clients=len({r['peer'] for r in sent if r['read']}),reply_messages=len(replies),replied_clients=len({r['peer'] for r in replies}),applications=sum(a['quality']=='valid' for a in apps),invalid_applications=sum(a['quality']!='valid' for a in apps),handoffs=len({r['thread_id'] for r in notices if r['status']=='delivered'}),handoffs_pending=len({r['thread_id'] for r in notices if r['status']=='pending'}),warming_now=len(cohorts),refusal_clients=len({r['peer'] for r in refused}),ai_requests=len(usage)),applications=apps,messages=sorted(details,key=lambda r:r['at']),refusals=refused,usage=usage,health=health,warnings=warnings,research_records=len(research),safety=safety,costs=costs,sync_pending=sync_pending)

LABELS={'sent_messages':'Отправлено сообщений','contacted_clients':'Написали клиентам','read_messages':'Прочитано сообщений, подтверждено','read_clients':'Клиентов с подтверждённым прочтением','reply_messages':'Входящих сообщений за период','replied_clients':'Клиентов с входящими за период','applications':'Заявок, действительных сейчас','invalid_applications':'Заявок с отменённым качеством','handoffs':'Клиентов передано МОП как заявка','handoffs_pending':'Ожидают передачи МОП как заявка','warming_now':'На прогреве сейчас','refusal_clients':'Клиентов с явной фразой отказа','ai_requests':'Записано запросов ИИ'}
def stamp(v):return datetime.fromtimestamp(v,TZ).strftime('%d.%m.%Y %H:%M') if v else 'нет данных'
def text(report):
 s=report['scope'];title=s['organization']+(' / '+s['company_name'] if s['company_name'] else '')
 out=['Отчёт: '+title,stamp(report['start'])+' — '+stamp(report['end']-1)+' (Дубай)',s['manager_id'],'Продукт: '+(report['product'] or 'все разрешённые')]+[LABELS[k]+': '+str(v) for k,v in report['metrics'].items()]
 h=report['health'];safety=report['safety'];out+=['','WORK: '+('подключена, сборщик отвечает' if h and report['generated_at']-h['updated_at']<=120 and h['connected'] else 'нет свежего подтверждения подключения'),'Клиентские отправки: '+('ОСТАНОВЛЕНЫ защитой с '+stamp(safety['created_at']) if safety['blocked'] else 'защитная остановка не установлена' if safety['known'] else 'состояние защиты не подтверждено'), 'Причина защиты: '+str(safety['operation'] or 'нет данных'), 'Сверка прочтений: '+stamp(h.get('receipts_at') if h else None),'Учтённые токены этого менеджера (НЕ общие списания): $'+format(report['costs']['usd_min'],'.4f')+'–$'+format(report['costs']['usd_max'],'.4f')+'; учёт неполный.', 'Учтённые вызовы инструментов: $'+format(report['costs'].get('tool_usd_known',0),'.4f')+'; возможны другие расходы.', 'Фактические списания общего ключа: смотрите сводку центра/OpenAI, здесь не определяются.', 'Запросов без расчёта цены: '+str(report['costs']['unknown_requests']), 'Без полного учёта инструментов: '+str(report['costs']['incomplete_tools']), 'Ожидает синхронизации записей: '+str(report['sync_pending']),'Прочтения — на сейчас; прогрев — на сейчас.','Отказы — по явным фразам; подробности — в Excel.']
 if not h or not h.get('receipts_at'):out+=['Прочтения пока неполные: идёт сверка.']
 return '\n'.join(out)

def excel(report):
 # This is an in-product exporter, using the application's installed dependency.
 from openpyxl import Workbook
 from openpyxl.styles import Font,PatternFill,Alignment
 from openpyxl.utils import get_column_letter
 wb=Workbook();wb.remove(wb.active)
 def sheet(name,headers,records):
  ws=wb.create_sheet(name);ws.append(headers)
  for record in records:
   row=[]
   for v in record:
    if isinstance(v,(dict,list)):v=json.dumps(v,ensure_ascii=False)
    if isinstance(v,str):
     v=''.join(ch for ch in v if ord(ch)>=32 or ch in '\n\t')[:32000]
     if v.lstrip().startswith(('=','+','-','@')):v="'"+v
    row.append(v)
   ws.append(row)
  ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
  for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='193B56')
  for cells in ws.iter_rows(min_row=2):
   for cell in cells:cell.alignment=Alignment(vertical='top',wrap_text=True)
  for i in range(1,len(headers)+1):ws.column_dimensions[get_column_letter(i)].width=55 if i==len(headers) else 25
  return ws
 sheet('Сводка',['Показатель','Значение'],[['Организация',report['scope']['organization']],['Менеджер',report['scope']['manager_id']],['Период Дубай',stamp(report['start'])+' — '+stamp(report['end']-1)],['Сформирован',stamp(report['generated_at'])],['Защитная остановка отправок',str(report['safety']['blocked'])],['Токены USD минимум (частичный учёт)',report['costs']['usd_min']],['Токены USD максимум (частичный учёт)',report['costs']['usd_max']],['Запросы без расчёта цены',report['costs']['unknown_requests']],['Неполный учёт инструментов',report['costs']['incomplete_tools']]]+[[LABELS[k],v] for k,v in report['metrics'].items()])
 sheet('Заявки',['ID','Клиент','Продукт','Потребность','Статус','Качество','Дата','Сообщение клиента','Источник','Уточнения','Основания'],[[a['id'],a['name'],a['product'],a['subject'],a['status'],a['quality'],stamp(a['created_at']),a['need_text'],a['url'],a.get('qualification','{}'),dict(events=a['events'],transfers=a['transfers'])] for a in report['applications']])
 sheet('Сообщения',['Направление','Клиент','Продукт','Дата','Подтверждено прочтение','Источник','Текст'],[[r['direction'],r['name'],r['product'],stamp(r['at']),('да' if r.get('read') else 'не подтверждено') if r['direction']=='Исходящее' else '',r['url'],r['text']] for r in report['messages']])
 sheet('Отказы',['Клиент','Дата','Источник','Фраза'],[[r['name'],stamp(r['at']),r['url'],r['text']] for r in report['refusals']])
 sheet('Расходы ИИ',['Этап','Модель','Статус','Дата','Входные токены','Выходные токены','USD минимум, токены','USD максимум, токены','Инструменты учтены'],[[u['phase'],u['model'],u['status'],stamp(u['created_at']),u['input_tokens'],u['output_tokens'],u['usd_min'],u['usd_max'],'да' if u['tools_complete'] else 'не полностью'] for u in report['usage']])
 sheet('Оговорки',['№','Описание'],[[i+1,w] for i,w in enumerate(report['warnings'])])
 stream=io.BytesIO();wb.save(stream);stream.seek(0);return stream
